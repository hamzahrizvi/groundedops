"""
Tests for parsing.py — document extraction quality.

The bugs being guarded against, all observed on real documents:

  1. python-docx keeps table cells in doc.tables, NOT doc.paragraphs, so
     reading only paragraphs dropped every table in every .docx silently.
     For product documentation that is most of the specifications.

  2. pypdf's extract_text() flattens a table's cells into one space-joined
     run, so "URL | BASE_URL/..." became "URL BASE_URL/..." and the
     label/value relationship was lost. A question whose answer sat in a spec
     table retrieved that table at 0.99 similarity and was still refused.

  3. Running headers/footers were indexed once per page — 88 copies of
     "Document Revision - v.15 ICU_Network_API - 10" in one manual.
"""

import os
import tempfile

from docx import Document

import parsing
from parsing import _render_table, _strip_repeated_lines, extract_pages


# ── _render_table ──────────────────────────────────────────────────────────

def test_render_table_two_columns_becomes_label_value():
    rows = [["URL", "BASE_URL/credentials"], ["Method", "POST"]]
    out = _render_table(rows)
    assert "URL: BASE_URL/credentials" in out
    assert "Method: POST" in out


def test_render_table_wide_keeps_cells_separated():
    rows = [["Status", "Body", "Notes"], ["200", "{}", "ok"]]
    out = _render_table(rows)
    # Pipe-separated, so a model can still see which cell is which column.
    assert "Status | Body | Notes" in out
    assert "200 | {} | ok" in out


def test_render_table_newlines_inside_cells_do_not_break_rows():
    rows = [["Headers", "Content-Type: application/json\nAuthorization: Bearer"]]
    out = _render_table(rows)
    assert out.count("\n") == 0, "a cell's newline must not split the row"
    assert "Authorization: Bearer" in out


def test_render_table_ignores_entirely_blank_rows():
    assert _render_table([["", ""], ["A", "B"]]) == "A: B"
    assert _render_table([]) == ""
    assert _render_table(None) == ""


def test_render_table_ragged_rows_do_not_raise():
    rows = [["a", "b", "c"], ["d"]]
    out = _render_table(rows)
    assert "a | b | c" in out


# ── .docx tables ───────────────────────────────────────────────────────────

def _docx_with_table(path):
    doc = Document()
    doc.add_paragraph("Specifications follow.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Operating voltage"
    table.cell(0, 1).text = "12VDC"
    table.cell(1, 0).text = "Protocols"
    table.cell(1, 1).text = "SSP, ccTalk, SI2"
    doc.add_paragraph("End of section.")
    doc.save(path)


def test_docx_table_content_is_extracted():
    tmp = os.path.join(tempfile.mkdtemp(), "spec.docx")
    _docx_with_table(tmp)

    pages = extract_pages(tmp)
    assert len(pages) == 1
    text = pages[0][1]

    # The whole point: this used to be absent entirely.
    assert "Operating voltage: 12VDC" in text
    assert "Protocols: SSP, ccTalk, SI2" in text
    # And the surrounding prose is still there.
    assert "Specifications follow." in text
    assert "End of section." in text


def test_docx_keeps_document_order():
    """A table must stay between the paragraphs that introduce and follow it —
    appending all tables at the end separates a spec from its context."""
    tmp = os.path.join(tempfile.mkdtemp(), "order.docx")
    _docx_with_table(tmp)

    text = extract_pages(tmp)[0][1]
    intro = text.index("Specifications follow.")
    spec = text.index("Operating voltage: 12VDC")
    end = text.index("End of section.")
    assert intro < spec < end


# ── running headers/footers ────────────────────────────────────────────────

def test_strip_repeated_lines_drops_furniture_keeps_content():
    footer = "Document Revision - v.15 ICU_Network_API - 10"
    pages = [(i, f"{footer}\nUnique content for page {i}.") for i in range(1, 9)]

    out = _strip_repeated_lines(pages)

    assert len(out) == 8
    for num, text in out:
        assert footer not in text
        assert f"Unique content for page {num}." in text


def test_strip_repeated_lines_leaves_short_documents_alone():
    """On a 3-page file a repeated line is as likely to be real content."""
    pages = [(i, "Shared line\nBody %d" % i) for i in range(1, 4)]
    assert _strip_repeated_lines(pages) == pages


def test_strip_repeated_lines_keeps_long_repeated_passages():
    """A long verbatim repeat is more likely a repeated warning than
    furniture, so the filter is length-bounded."""
    warning = ("W" * 200)
    pages = [(i, f"{warning}\nBody {i}") for i in range(1, 9)]
    out = _strip_repeated_lines(pages)
    assert all(warning in text for _, text in out)


# ── pages with no extractable text ─────────────────────────────────────────

def test_unsupported_extension_returns_empty_not_raise():
    tmp = os.path.join(tempfile.mkdtemp(), "thing.rtf")
    with open(tmp, "w", encoding="utf-8") as fh:
        fh.write("hello")
    assert extract_pages(tmp) == []


def test_missing_file_returns_empty_not_raise():
    assert extract_pages(os.path.join(tempfile.mkdtemp(), "nope.pdf")) == []


def test_txt_is_one_page():
    tmp = os.path.join(tempfile.mkdtemp(), "notes.txt")
    with open(tmp, "w", encoding="utf-8") as fh:
        fh.write("plain text body")
    assert extract_pages(tmp) == [(1, "plain text body")]


def test_pdfplumber_is_the_active_reader():
    """If pdfplumber is missing the code still works, but tables get flattened
    by pypdf -- which is the failure mode this module exists to fix. Assert the
    install is present so a dropped dependency is loud, not silent."""
    assert parsing.pdfplumber is not None, "pdfplumber not installed"


# ── headings vs wrapped prose ──────────────────────────────────────────────
# Bug 4: the opening paragraph of the NV4000 Development Kit is set larger
# than body text, so both of its wrapped lines passed _is_heading -- short
# enough, alphabetic, no terminal full stop. The first became a section
# heading with the second as the next heading, leaving it with an empty body
# that dropped out at ingest. The phrase "multi note recycler and bill
# validator" was in the document and absent from the index, so the one
# question every product gets asked -- what is it? -- could not be answered
# from the document that says.

def test_line_starting_lower_case_is_prose_not_a_heading():
    assert parsing._continues_previous(
        "validator. To help you with your integration we have provided")


def test_line_spanning_a_sentence_break_is_prose_not_a_heading():
    assert parsing._continues_previous("Ready to run. Connect the unit")


def test_real_headings_are_not_mistaken_for_prose():
    for heading in ("Currency Datasets", "3D CAD Files", "Contact Us",
                    "Software Development Kit", "NV4000 Manual"):
        assert not parsing._continues_previous(heading), heading


def test_abbreviation_mid_line_is_not_a_sentence_break():
    # "Fig. 3" and "e.g. the bezel" are followed by a digit or lower case,
    # so requiring a capital after the stop leaves them alone.
    assert not parsing._continues_previous("Fig. 3 shows the bezel")


# ── tables keep their own heading ──────────────────────────────────────────
# Bug 5: every table was appended after ALL of the page's prose, so each one
# landed under the LAST heading on its page. The 3D CAD table was filed under
# "Software Development Kit"; "3D CAD Files" kept nothing and never reached
# the index; and a question about the CAD files could not find the rows that
# answer it while one naming a part number could.

class _FakeTable:
    def __init__(self, top, rows):
        self.bbox = (0, top, 500, top + 40)
        self._rows = rows

    def extract(self):
        return self._rows


class _FakePage:
    """Enough of a pdfplumber page for _page_blocks: text lines with tops."""

    def __init__(self, lines):
        self._lines = lines

    def extract_text_lines(self):
        return [{"text": t, "top": top} for t, top in self._lines]


def test_each_table_stays_under_its_own_heading():
    page = _FakePage([("3D CAD Files", 100), ("Software Development Kit", 300)])
    prose = "3D CAD Files\nSoftware Development Kit"
    tables = [_FakeTable(200, [["NV4000-00000", "1000 Note Cashbox"]]),
              _FakeTable(400, [["ITL SDK Package Manual", "REST API endpoints"]])]

    blocks = parsing._page_blocks(page, prose, tables)
    kinds = [k for k, _ in blocks]
    assert kinds == ["prose", "table", "prose", "table"], blocks
    assert "3D CAD Files" in blocks[0][1]
    assert "NV4000-00000" in blocks[1][1]
    assert "Software Development Kit" in blocks[2][1]
    assert "ITL SDK Package Manual" in blocks[3][1]


def test_page_blocks_falls_back_when_lines_do_not_line_up():
    # One line object, two lines of text: the split cannot be trusted, so the
    # old order is kept rather than risking a dropped or duplicated line.
    page = _FakePage([("Heading", 100)])
    blocks = parsing._page_blocks(page, "Heading\nand more prose",
                                  [_FakeTable(50, [["a", "b"]])])
    assert [k for k, _ in blocks] == ["prose", "table"]


def test_page_blocks_with_no_tables_is_one_prose_block():
    page = _FakePage([("Contact Us", 100)])
    blocks = parsing._page_blocks(page, "Contact Us", [])
    assert blocks == [("prose", "Contact Us")]
