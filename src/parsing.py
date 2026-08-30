import logging

from pypdf import PdfReader
from docx import Document

logger = logging.getLogger(__name__)

# pdfplumber is the layout-aware reader (see extract_pages). Imported lazily-ish
# here so a missing install degrades to pypdf instead of breaking ingestion
# outright — the old behaviour, which is worse but not nothing.
#
# Chosen over PyMuPDF deliberately: PyMuPDF is AGPL-3.0 unless commercially
# licensed, and this ships as a commercial product. pdfplumber is MIT.
try:
    import pdfplumber
except Exception as _exc:                                    # pragma: no cover
    pdfplumber = None
    logger.warning(f"pdfplumber unavailable ({_exc}); PDF tables will be "
                   f"flattened by pypdf and may be unusable")


# Sentinels marking an extracted table so the chunker can treat it as one
# atomic unit. Deliberately ugly and unlikely to occur in a manual. Stripped
# by chunking.strip_table_fences() before anything is indexed.
TABLE_OPEN = "<<<GO_TABLE>>>"
TABLE_CLOSE = "<<</GO_TABLE>>>"


def _render_table(rows: list[list[str | None]] | None) -> str:
    """A table as text a language model can actually read.

    pypdf's extract_text() flattens a table into its cells separated by
    spaces, so "URL | BASE_URL/..." arrives as "URL BASE_URL/..." and the
    label/value relationship is gone. Observed consequence: a question whose
    answer sat in a spec table retrieved that table at 0.99 similarity and
    was still refused, because the model could not tell what applied to what.

    Two-column tables become "label: value" lines, which is how a spec sheet
    reads anyway. Wider tables keep a header row and pipe-separated cells.
    """
    clean = []
    for row in rows or []:
        cells = [(c or "").replace("\n", " ").strip() for c in row]
        if any(cells):
            clean.append(cells)
    if not clean:
        return ""

    width = max(len(r) for r in clean)

    if width == 2:
        out = []
        for cells in clean:
            label, value = cells[0], cells[1] if len(cells) > 1 else ""
            if label and value:
                out.append(f"{label}: {value}")
            else:
                out.append(label or value)
        return "\n".join(out)

    lines = []
    for cells in clean:
        cells = cells + [""] * (width - len(cells))
        lines.append(" | ".join(cells).strip())
    return "\n".join(lines)


def _pdf_pages_plumber(path: str) -> tuple[list[tuple[int, str]], list[int]]:
    """Layout-aware PDF text, with tables rendered separately.

    Table regions are excluded from the prose pass and re-emitted through
    _render_table, so the structured form replaces the flattened one rather
    than sitting alongside it — duplicating both would inflate the index and
    skew BM25 term counts.

    Returns (pages, empty_page_numbers).
    """
    out: list[tuple[int, str]] = []
    empty: list[int] = []

    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            try:
                tables = page.find_tables()
            except Exception:
                tables = []

            boxes = [t.bbox for t in tables]

            def outside_tables(obj, _boxes=boxes):
                # Words are kept only if they fall outside every table bbox.
                for x0, top, x1, bottom in _boxes:
                    if (obj.get("x0", 0) >= x0 and obj.get("x1", 0) <= x1
                            and obj.get("top", 0) >= top
                            and obj.get("bottom", 0) <= bottom):
                        return False
                return True

            try:
                prose = (page.filter(outside_tables).extract_text() or "") if boxes \
                    else (page.extract_text() or "")
            except Exception:
                prose = page.extract_text() or ""

            parts = [prose.strip()] if prose.strip() else []
            for t in tables:
                try:
                    rendered = _render_table(t.extract())
                except Exception:
                    rendered = ""
                if rendered.strip():
                    # Fenced so the chunker can keep a table whole. A spec
                    # table split mid-row loses the row: "Validator NV9S" in
                    # one chunk and "1.05 Kg" in the next answers nothing,
                    # and it is exactly this corpus's most-asked content.
                    # chunking.strip_table_fences() removes these before the
                    # text is stored, so the marker never reaches the index
                    # or a prompt.
                    parts.append(TABLE_OPEN + "\n" + rendered.strip()
                                 + "\n" + TABLE_CLOSE)

            text = "\n\n".join(parts).strip()
            if text:
                out.append((i, text))
            else:
                empty.append(i)

    return out, empty


def _pdf_pages_pypdf(path: str) -> tuple[list[tuple[int, str]], list[int]]:
    """Fallback: plain pypdf extraction, no layout model."""
    out: list[tuple[int, str]] = []
    empty: list[int] = []
    reader = PdfReader(path)
    for i, page in enumerate(reader.pages, start=1):
        txt = page.extract_text() or ""
        if txt.strip():
            out.append((i, txt))
        else:
            empty.append(i)
    return out, empty


def _docx_text(path: str) -> str:
    """Paragraphs AND tables, in document order.

    python-docx keeps table cells in doc.tables, NOT in doc.paragraphs — so
    reading only paragraphs dropped every table in every .docx silently. For
    product documentation that is most of the specifications. Walking the body
    XML keeps tables positioned where they actually appear, rather than
    appending them all at the end away from the text that introduces them.
    """
    doc = Document(path)
    body = doc.element.body
    para_by_el = {p._element: p for p in doc.paragraphs}
    table_by_el = {t._element: t for t in doc.tables}

    parts: list[str] = []
    for child in body.iterchildren():
        if child in para_by_el:
            txt = para_by_el[child].text.strip()
            if txt:
                parts.append(txt)
        elif child in table_by_el:
            rows = [[cell.text for cell in row.cells]
                    for row in table_by_el[child].rows]
            rendered = _render_table(rows)
            if rendered.strip():
                parts.append(rendered.strip())

    return "\n".join(parts)


def _strip_repeated_lines(pages: list[tuple[int, str]],
                          threshold: float = 0.6) -> list[tuple[int, str]]:
    """Drop running headers/footers.

    A line appearing on most pages of a document is furniture, not content:
    "Document Revision - v.15 ICU_Network_API - 10" was being indexed 88 times
    in one manual. That inflates BM25 term counts, and it consumes chunk space
    that should hold the answer.

    Only applied to documents long enough for the signal to mean something —
    on a three-page file a genuinely repeated line is as likely to be real.
    """
    if len(pages) < 4:
        return pages

    from collections import Counter
    counts: Counter = Counter()
    for _, text in pages:
        for line in {ln.strip() for ln in text.splitlines() if ln.strip()}:
            # Structural markers are on their own line and repeat on every
            # page that has a table, so they looked exactly like a running
            # header and were deleted -- which silently turned OFF
            # table-aware chunking for any document whose tables appear on
            # more than 60% of its pages. Measured before this guard: 0 of
            # 44 tables fenced in MyCheckr Mini, 0 of 140 in ICU_Network_API,
            # 0 of 57 in MyCheckr v7, while NV9 and SMART Coin (tables on
            # fewer pages) came through fine. That is exactly the shape of
            # bug that reads as "it works for some products, not others".
            if line in (TABLE_OPEN, TABLE_CLOSE):
                continue
            counts[line] += 1

    cutoff = max(2, int(len(pages) * threshold))
    # Short lines only: a long paragraph repeated verbatim is more likely to be
    # real duplicated content (a repeated warning, say) than page furniture.
    furniture = {ln for ln, n in counts.items() if n >= cutoff and len(ln) <= 120}
    if not furniture:
        return pages

    logger.info("Dropped %d repeated header/footer line(s) across %d pages",
                len(furniture), len(pages))
    out = []
    for num, text in pages:
        kept = "\n".join(ln for ln in text.splitlines()
                         if ln.strip() not in furniture)
        if kept.strip():
            out.append((num, kept))
    return out


def extract_pages(path: str) -> list[tuple[int, str]]:
    """Extract text as [(page_number, text), ...], 1-indexed.

    v12.0. extract_text() below joined every page into one string, which
    destroyed page boundaries before chunking ever ran — so a chunk could
    never say which page it came from, and citations couldn't reference
    one. This preserves the boundary so ingest can tag each chunk with its
    page span.

    Non-paginated formats report a single page 1: .docx has no fixed
    pagination without rendering (page breaks depend on the renderer), and
    .txt has none at all. Treating them as one page is honest — better
    than inventing page numbers that wouldn't match what the reader sees.

    Pages that yield no text are counted and reported. They used to be
    skipped in silence, so a scanned or image-only page simply was not in the
    knowledge base and nobody was told — indistinguishable from a page whose
    content the model just could not use.
    """
    try:
        if path.endswith(".txt"):
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return [(1, f.read())]

        elif path.endswith(".pdf"):
            if pdfplumber is not None:
                try:
                    out, empty = _pdf_pages_plumber(path)
                except Exception as exc:
                    logger.warning(f"Layout-aware extraction failed for "
                                   f"'{path}' ({exc}); falling back to pypdf")
                    out, empty = _pdf_pages_pypdf(path)
            else:
                out, empty = _pdf_pages_pypdf(path)

            if empty:
                logger.warning(
                    "%s: %d of %d page(s) produced no text (%s). These are "
                    "most likely scanned or image-only and are NOT searchable "
                    "— they need OCR to be usable.",
                    path, len(empty), len(empty) + len(out),
                    ", ".join(str(p) for p in empty[:12])
                    + (" ..." if len(empty) > 12 else ""))
            return _strip_repeated_lines(out)

        elif path.endswith(".docx"):
            txt = _docx_text(path)
            return [(1, txt)] if txt.strip() else []

        else:
            logger.warning(f"Unsupported file type: {path}")
            return []

    except Exception as exc:
        logger.error(f"Failed to extract pages from '{path}': {exc}")
        return []


def extract_text(path: str) -> str:
    """Whole-document text. Retained for callers that don't need pages
    (eval, diagnostics); ingest now uses extract_pages()."""
    return "\n\n".join(t for _, t in extract_pages(path))
